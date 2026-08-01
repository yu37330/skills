# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""High-level API for sdpm — single entry points for generate, measure, preview, init, code_block.

These functions encapsulate the full workflow that the CLI (pptx_builder.py) performs.
mcp-local and other consumers should call these instead of assembling low-level APIs.
"""

import os
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any


def get_templates_dirs() -> list[Path]:
    """Return ordered list of directories to search for bundled/user-local templates.

    Search order (first match wins):
      1. $SDPM_TEMPLATES_DIR — os.pathsep-separated list (same semantics as PATH)
      2. get_user_config_dir()/templates/ — user-local templates
      3. Package-bundled templates/ directory (skill/templates/)
    """
    from sdpm.config import _get_resource_dirs

    bundled = Path(__file__).parent.parent / "templates"
    return _get_resource_dirs("SDPM_TEMPLATES_DIR", "templates", bundled)


def get_styles_dirs() -> list[Path]:
    """Return ordered list of directories to search for style HTMLs.

    Search order (first match wins):
      1. $SDPM_STYLES_DIR — os.pathsep-separated list (same semantics as PATH)
      2. get_user_config_dir()/styles/ — user-local styles
      3. Package-bundled references/examples/styles/ directory
    """
    from sdpm.config import _get_resource_dirs
    from sdpm.reference import BUNDLED_STYLES_DIR

    return _get_resource_dirs("SDPM_STYLES_DIR", "styles", BUNDLED_STYLES_DIR)


def list_styles_filtered(
    styles_dirs: list[Path],
    pinned_names: list[str],
    include_all: bool = False,
) -> list[dict]:
    """List styles with pin/source metadata, optionally filtered.

    Filesystem-based entry point for MCP Local / CLI.
    Determines source ("user" vs "builtin") by checking whether each style
    lives in the user-local directory.

    Args:
        styles_dirs: Ordered directories from get_styles_dirs().
        pinned_names: Pinned style names from state.json.
        include_all: Pass through to filter_styles().

    Returns:
        Filtered list with pinned/source metadata.
    """
    from sdpm.config import get_user_config_dir
    from sdpm.reference import filter_styles, list_styles_merged

    user_dir = get_user_config_dir() / "styles"
    raw = list_styles_merged(styles_dirs)

    # Tag source based on whether the style file exists in user dir
    for s in raw:
        if (user_dir / f"{s['name']}.html").exists():
            s["source"] = "user"
        else:
            s["source"] = "builtin"

    return filter_styles(raw, pinned_names, include_all)


def list_templates_with_metadata(
    templates_dirs: list[Path],
    metadata: dict[str, dict],
) -> list[dict]:
    """List templates with source and metadata.

    Pure function — no I/O beyond filesystem glob.

    Consumers are the local paths only (L1 CLI, mcp-local, local Web UI via
    state.json). The cloud path (mcp-server) does NOT use this function — it has
    an independent listing in mcp-server/tools/template.py where builtin notes
    are first-class DDB items (BUILTIN_NOTE#<name>), merged at the tool layer.
    Keep that divergence in mind before consolidating.

    Args:
        templates_dirs: From get_templates_dirs(). Last entry is bundled.
        metadata: {name: {description, theme_colors, fonts, layout_count}} from state.json.
            Builtin templates are looked up under "builtin:<name>" first (the key
            local Web UI uses for its analysis cache; user notes share the entry),
            falling back to "<name>" for backward compatibility.

    Returns:
        Sorted list of template dicts with name, source, description, theme_colors, fonts, layout_count.
    """
    bundled_dir = templates_dirs[-1] if templates_dirs else None
    seen: dict[str, dict] = {}
    for d in templates_dirs:
        if not d.exists():
            continue
        for t in sorted(d.glob("*.pptx")):
            name = t.stem
            if name in seen:
                continue
            source = "builtin" if d == bundled_dir else "user"
            meta = metadata.get(name, {})
            if source == "builtin":
                # builtin: prefix keeps builtin metadata (incl. user notes) from
                # colliding with a same-named user template entry.
                meta = metadata.get(f"builtin:{name}") or meta
            seen[name] = {
                "name": name,
                "source": source,
                "description": meta.get("description", ""),
                "theme_colors": meta.get("theme_colors", {}),
                "fonts": meta.get("fonts", {}),
                "layout_count": meta.get("layout_count", 0),
            }
    return sorted(seen.values(), key=lambda x: (x["source"] != "user", x["name"]))


def analyze_and_store_template(template_path: Path, description: str = "") -> dict:
    """Analyze a template and return metadata for storage.

    Calls the existing analyze_template() and reshapes the result.
    Persistence is the caller's responsibility (state.json or DDB).

    Args:
        template_path: Path to .pptx file.
        description: User-provided description.

    Returns:
        Dict with name, description, theme_colors, fonts, layout_count, layouts.
    """
    from sdpm.analyzer import analyze_template as _analyze

    result = _analyze(template_path)
    return {
        "name": template_path.stem,
        "description": description,
        "theme_colors": result.get("theme_colors", {}),
        "fonts": result.get("fonts", {}),
        "layout_count": len(result.get("layouts", [])),
        "layouts": result.get("layouts", []),
    }


def apply_style(deck_dir: str | Path, style: str) -> dict[str, Any]:
    """Apply a named style to a deck's art-direction.

    Copies the style HTML to {deck_dir}/specs/art-direction.html.

    Args:
        deck_dir: Deck output directory path.
        style: Style name (e.g. "elegant-dark").

    Returns:
        Dict with status, path, style. Or error key if not found.
    """
    import shutil

    styles_dirs = get_styles_dirs()
    src = _find_style_in_dirs(style, styles_dirs)
    if src is None:
        available = [p.stem for d in styles_dirs if d.is_dir() for p in d.glob("*.html")]
        return {"error": f"Style not found: {style}. Available: {sorted(set(available))}"}
    deck_path = Path(deck_dir)
    if not deck_path.is_dir():
        return {"error": f"Deck directory not found: {deck_dir}"}
    dest = deck_path / "specs" / "art-direction.html"
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest)
    return {"status": "ok", "path": str(dest), "style": style}


def _find_style_in_dirs(name: str, styles_dirs: list[Path]) -> Path | None:
    """Search for a style HTML by name across the given directories.

    Returns the first existing path, or None if not found.
    """
    filename = name if name.endswith(".html") else name + ".html"
    for d in styles_dirs:
        candidate = d / filename
        if candidate.exists():
            return candidate
    return None


def _find_template_in_dirs(name: str, templates_dirs: list[Path]) -> Path | None:
    """Search for a template by name across the given directories.

    Returns the first existing path, or None if not found.
    """
    filename = name if name.endswith(".pptx") else name + ".pptx"
    for d in templates_dirs:
        candidate = d / filename
        if candidate.exists():
            return candidate
    return None


def _resolve_template(
    data: dict,
    input_path: str | Path | None,
    templates_dirs: list[Path],
) -> tuple[Path, bool]:
    """Resolve template path from presentation data.

    Returns (template_path, custom_template) or raises FileNotFoundError.
    """
    if data.get("template"):
        if input_path:
            p = Path(input_path)
            base_dir = p if p.is_dir() else p.parent
        else:
            base_dir = Path(".")
        template = base_dir / data["template"]
        if template.exists():
            return template, True
        found = _find_template_in_dirs(data["template"], templates_dirs)
        if found is not None:
            return found, True
    raise FileNotFoundError('No template specified. Set "template" in presentation JSON.')


def _get_output_base_dir() -> Path:
    """Get output base directory from config, with WSL fallback."""
    env_dir = os.environ.get("SDPM_OUTPUT_DIR")
    if env_dir:
        return Path(env_dir)
    try:
        from sdpm.config import get_output_dir

        return get_output_dir()
    except Exception:
        pass
    from sdpm.preview.backend import _is_wsl

    if _is_wsl():
        import subprocess

        try:
            result = subprocess.run(  # nosec B603 # nosemgrep: python.lang.security.audit.dangerous-subprocess-use-audit
                [
                    "/mnt/c/Windows/System32/WindowsPowerShell/v1.0/powershell.exe",
                    "-Command",
                    "[Environment]::GetFolderPath('MyDocuments')",
                ],
                capture_output=True,
                timeout=10,
            )
            win_path = result.stdout.decode("cp932", errors="replace").strip()
            if win_path:
                wsl = subprocess.run(["wslpath", win_path], capture_output=True, text=True)  # nosec B603 # nosemgrep: python.lang.security.audit.dangerous-subprocess-use-audit
                if wsl.returncode == 0:
                    return Path(wsl.stdout.strip()) / "SDPM-Presentations"
        except Exception:
            pass
    return Path.home() / "Documents" / "SDPM-Presentations"


def init(
    name: str,
    template: str | Path | None = None,
    output_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Initialize a presentation workspace.

    Creates output directory with deck.json, slides/, and specs/.

    Args:
        name: Presentation name (used in directory name).
        template: Template name or path. If provided, extracts fonts.
        output_dir: Explicit output directory. Auto-generated if None.

    Returns:
        Dict with output_dir, deck_json, template, fonts, workspace.
    """
    from sdpm.analyzer import extract_fonts
    from sdpm.utils.io import write_json

    if output_dir:
        out_dir = Path(output_dir).expanduser()
    else:
        ts = datetime.now().strftime("%Y%m%d-%H%M")
        dir_name = f"{ts}-{name}" if name else ts
        out_dir = _get_output_base_dir() / dir_name
    out_dir.mkdir(parents=True, exist_ok=True)

    deck_data: dict[str, Any] = {
        "template": "",
        "fonts": {"fullwidth": "", "halfwidth": ""},
        "defaultTextColor": "",
    }

    if template:
        template_src = Path(template).expanduser()
        if not template_src.exists():
            found = _find_template_in_dirs(str(template), get_templates_dirs())
            if found is not None:
                template_src = found
        if template_src.exists():
            template_src = template_src.resolve()
            deck_data["template"] = template_src.name
            try:
                deck_data["fonts"] = extract_fonts(template_src)
            except Exception:
                pass

    deck_json = out_dir / "deck.json"
    write_json(deck_json, deck_data, suffix="\n")

    (out_dir / "slides").mkdir(exist_ok=True)
    specs_dir = out_dir / "specs"
    specs_dir.mkdir(exist_ok=True)
    spec_files = ("brief.md", "outline.md")
    for spec_name in spec_files:
        (specs_dir / spec_name).touch()

    return {
        "output_dir": str(out_dir),
        "deck_json": str(deck_json),
        "template": deck_data.get("template", ""),
        "fonts": deck_data.get("fonts", {}),
        "workspace": ["deck.json", "slides/"] + [f"specs/{s}" for s in spec_files],
    }


@dataclass
class BuildConfig:
    """Resolved configuration for PPTX build."""

    template_path: Path
    custom_template: bool
    fonts: dict
    default_text_color: str
    slides: list[dict] = field(default_factory=list)  # override解決済み
    base_dir: Path = field(default_factory=lambda: Path("."))
    warnings: list[str] = field(default_factory=list)
    lint_diagnostics: list = field(default_factory=list)
    auto_spacing: bool = True  # deck.json "autoSpacing"; False for imported decks


def _assemble_slides_from_dir(
    deck_dir: Path,
    only_slugs: set[str] | None = None,
) -> tuple[dict, list[dict]]:
    """Assemble presentation dict from deck.json + outline.md + slides/*.json.

    Args:
        deck_dir: Directory containing deck.json, specs/outline.md, slides/*.
        only_slugs: When set, include only the specified slugs (preserving
            outline order). None means include all.

    Returns:
        (deck_meta, slides) where deck_meta has template/fonts/defaultTextColor
        and slides is the ordered list from outline.md.
    """
    from sdpm.utils.io import read_json

    deck_json = deck_dir / "deck.json"
    if not deck_json.exists():
        raise FileNotFoundError(f"deck.json not found in {deck_dir}")
    deck_meta = read_json(deck_json)

    slugs = parse_outline_slugs(deck_dir / "specs" / "outline.md")
    if only_slugs is not None:
        slugs = [s for s in slugs if s in only_slugs]

    slides: list[dict] = []
    for slug in slugs:
        slide_path = deck_dir / "slides" / f"{slug}.json"
        if not slide_path.exists():
            continue  # skip missing slides silently
        slide = read_json(slide_path)
        slide.setdefault("id", slug)
        slides.append(slide)

    return deck_meta, slides


def parse_outline_slugs(outline_path: Path) -> list[str]:
    """Parse outline.md and return ordered list of slugs.

    Format: ``- [slug] Message text``

    Args:
        outline_path: Path to outline.md.

    Returns:
        List of slug strings in document order.
    """
    import re

    pattern = re.compile(r"^-\s*\[([a-z0-9-]+)\]\s*")
    slugs: list[str] = []
    if not outline_path.exists():
        return slugs
    for line in outline_path.read_text(encoding="utf-8").splitlines():
        m = pattern.match(line)
        if m:
            slugs.append(m.group(1))
    return slugs


def _resolve_config(
    json_path: str | Path,
    only_slugs: set[str] | None = None,
) -> BuildConfig:
    """Resolve template, fonts, icons, overrides from JSON or directory.

    Accepts either a presentation.json file path (legacy) or a directory
    containing deck.json + specs/outline.md + slides/*.json (new format).

    Args:
        json_path: Path to presentation.json or deck directory.
        only_slugs: When set, include only the specified slugs in the build.

    Raises FileNotFoundError, ValueError on missing template/icons.
    """
    from sdpm.builder import resolve_override, validate_icons_in_json
    from sdpm.utils.io import read_json

    input_path = Path(json_path)
    if not input_path.exists():
        raise FileNotFoundError(f"Not found: {json_path}")

    # Directory input: deck.json + slides/*.json
    if input_path.is_dir():
        deck_meta, slides = _assemble_slides_from_dir(input_path, only_slugs=only_slugs)
        data = {**deck_meta, "slides": slides}
        base_dir = input_path
    else:
        data = read_json(input_path)
        base_dir = input_path.parent

    templates_dirs = get_templates_dirs()
    warnings: list[str] = []

    template_file, custom = _resolve_template(data, str(input_path), templates_dirs)

    # Auto-fill fonts
    from sdpm.analyzer import extract_fonts as _extract_fonts
    from sdpm.analyzer import _extract_theme_colors_raw

    fonts = data.get("fonts")
    if not fonts or not fonts.get("fullwidth"):
        fonts = _extract_fonts(template_file)
        warnings.append("fonts auto-detected from template")

    # Auto-fill defaultTextColor
    dtc = data.get("defaultTextColor")
    if not dtc:
        _, is_dark = _extract_theme_colors_raw(template_file)
        dtc = "#FFFFFF" if is_dark else "#333333"
        warnings.append(f"defaultTextColor auto-set to {dtc}")

    # Lint
    from sdpm.schema.lint import lint as lint_slides

    lint_diagnostics = lint_slides(data)

    # Validate icons
    missing = validate_icons_in_json(data)
    if missing:
        raise ValueError(f"Missing assets ({len(missing)}): {', '.join(sorted(missing)[:10])}")

    # Token discipline: fontSize must come from --fs-* tokens in active style
    from sdpm.checks import check_font_size_tokens, check_includes, check_overlay_textbox

    fs_warnings = check_font_size_tokens(data, input_path)
    warnings.extend(fs_warnings)

    # Overlay textbox: textbox stacked on a shape with the same bounding box
    overlay_warnings = check_overlay_textbox(data)
    warnings.extend(overlay_warnings)

    # Include references: a missing/empty include silently drops its content.
    include_warnings = check_includes(data, base_dir)
    warnings.extend(include_warnings)

    # Resolve overrides
    slides = data.get("slides", [])
    id_map = {}
    for s in slides:
        if "id" in s:
            id_map[s["id"]] = s
    resolved_slides = [resolve_override(s, id_map) for s in slides]

    # Filter by only_slugs for legacy JSON path (directory path already filtered)
    if only_slugs is not None and not input_path.is_dir():
        resolved_slides = [s for s in resolved_slides if s.get("id") in only_slugs]

    return BuildConfig(
        template_path=template_file,
        custom_template=custom,
        fonts=fonts,
        default_text_color=dtc,
        slides=resolved_slides,
        base_dir=base_dir,
        warnings=warnings,
        lint_diagnostics=lint_diagnostics,
        auto_spacing=data.get("autoSpacing", True),
    )


def _build(config: BuildConfig, output_path: Path) -> Path:
    """Build PPTX from resolved config. Returns output path."""
    from sdpm.builder import PPTXBuilder

    builder = PPTXBuilder(
        config.template_path,
        custom_template=config.custom_template,
        fonts=config.fonts,
        base_dir=config.base_dir,
        default_text_color=config.default_text_color,
        auto_spacing=config.auto_spacing,
    )
    for s in config.slides:
        builder.add_slide(s)
    builder.save(output_path)
    return output_path


def generate(
    json_path: str | Path,
    output_path: str | Path | None = None,
    only_slugs: set[str] | None = None,
) -> dict[str, Any]:
    """Generate PPTX from JSON.

    Includes: template resolution, icon validation, build, check_layout_imbalance.

    Args:
        json_path: Path to the slides JSON file.
        output_path: Output .pptx path. Auto-generated if None.
        only_slugs: When set, build only the specified slugs (for iso.pptx).

    Returns:
        Dict with output_path, slide_count, slides summary, warnings.
    """
    from sdpm.preview import check_layout_imbalance_data

    config = _resolve_config(json_path, only_slugs=only_slugs)

    # Output path
    input_path = Path(json_path)
    if not output_path:
        ts = datetime.now().strftime("%Y%m%d-%H%M%S")
        p = input_path.with_suffix(".pptx")
        out = p.with_stem(f"{p.stem}_{ts}")
    else:
        out = Path(output_path)

    _build(config, out)

    imbalance = check_layout_imbalance_data(out, config.slides)
    if imbalance:
        for a in imbalance:
            config.warnings.append(f"page{a['slide']:02d} ({a['layout']}) offset: {a['offset']} ({a['direction']})")

    # Summary
    summary = []
    for i, s in enumerate(config.slides, 1):
        title = s.get("title", "(no title)")
        if isinstance(title, dict):
            title = title.get("text", "(no title)")
        summary.append(f"page{i:02d} - {title}")

    result: dict[str, Any] = {
        "output_path": str(out),
        "slide_count": len(config.slides),
        "slides": summary,
        "warnings": config.warnings,
    }
    if config.lint_diagnostics:
        result["errors"] = {"lintDiagnostics": config.lint_diagnostics}
    return result


def measure(
    json_path: str | Path,
    slides: list[int] | list[str] | None = None,
) -> str:
    """Build PPTX from JSON, convert to SVG, extract text bboxes.

    Args:
        json_path: Path to the slides JSON file or directory.
        slides: Slide numbers (1-based int) or slugs (str) to measure. None for all.

    Returns:
        Text report of bbox measurements.
    """
    import tempfile

    from sdpm.preview.backend import LibreOfficeBackend, get_work_dir
    from sdpm.preview.measure import format_measure_report, measure_from_svg

    config = _resolve_config(json_path)

    # Resolve slugs to page numbers and build reverse mapping
    page_to_slug: dict[int, str] | None = None
    slide_indices: list[int] | None = None

    if slides and isinstance(slides[0], str):
        slug_to_page = {}
        for i, s in enumerate(config.slides):
            sid = s.get("id", "")
            if sid:
                slug_to_page[sid] = i + 1
        slide_indices = [slug_to_page[slug] for slug in slides if slug in slug_to_page]
        page_to_slug = {v: k for k, v in slug_to_page.items()}
    else:
        slide_indices = slides

    input_path = Path(json_path)
    work_dir = get_work_dir(input_path.parent if input_path.is_dir() else input_path.resolve().parent)

    with tempfile.TemporaryDirectory(dir=work_dir) as tmp_dir:
        tmp_pptx = Path(tmp_dir) / "measure.pptx"
        _build(config, tmp_pptx)

        backend = LibreOfficeBackend()
        svg_path = backend.export_svg(tmp_pptx, work_dir=work_dir)
        if svg_path is None:
            raise RuntimeError("SVG export failed. Is LibreOffice (soffice) installed?")

        try:
            results = measure_from_svg(svg_path, slide_indices)
            return format_measure_report(results, page_to_slug=page_to_slug)
        finally:
            import shutil
            if svg_path:
                shutil.rmtree(svg_path.parent, ignore_errors=True)


def preview(
    json_path: str | Path,
    output_path: str | Path | None = None,
    pages: list[int] | None = None,
    grid: bool = False,
) -> dict[str, Any]:
    """Build PPTX from JSON and export slides as PNG images.

    Args:
        json_path: Path to the slides JSON file.
        output_path: Output .pptx path. Auto-generated if None.
        pages: Page numbers to export. None for all.
        grid: Add grid overlay to PNGs.

    Returns:
        Dict with preview_dir, files list, and output_path.
    """
    import glob
    import re
    import subprocess

    from pptx import Presentation

    from sdpm.preview import export_pdf

    config = _resolve_config(json_path)

    # Output path
    input_path = Path(json_path)
    if not output_path:
        ts = datetime.now().strftime("%Y%m%d-%H%M%S")
        p = input_path.with_suffix(".pptx")
        out = p.with_stem(f"{p.stem}_{ts}")
    else:
        out = Path(output_path)

    _build(config, out)

    # Preview dir — use project-local _work/ to avoid macOS EDR/DLP blocking system temp
    from sdpm.preview import get_work_dir
    work_dir = get_work_dir(input_path.parent if input_path.is_dir() else input_path.resolve().parent)
    out_dir = Path(tempfile.mkdtemp(dir=work_dir))

    pages_set = set(pages) if pages else None

    # PDF + pdftoppm pipeline
    pdf = out_dir / "slides.pdf"
    if not export_pdf(out, pdf, work_dir=work_dir):
        raise RuntimeError("PDF export failed. Is LibreOffice (soffice) installed?")

    cmd = ["pdftoppm", "-png", "-scale-to", "1280", str(pdf), str(out_dir / "page")]
    result = subprocess.run(cmd, capture_output=True, text=True, stdin=subprocess.DEVNULL)  # nosec B603 # nosemgrep: python.lang.security.audit.dangerous-subprocess-use-audit
    if result.returncode != 0:
        raise RuntimeError(f"PNG conversion failed. Is poppler (pdftoppm) installed? {result.stderr}")

    # Rename with slide titles
    prs = Presentation(str(out))
    titles = _extract_slide_titles(prs)

    generated = []
    for png in sorted(glob.glob(str(out_dir / "page-*.png"))):
        match = re.match(r"page-(\d+)\.png", Path(png).name)
        if match:
            num = int(match.group(1))
            if pages_set and num not in pages_set:
                Path(png).unlink()
                continue
            new_name = f"page{num:02d}-{titles.get(num, 'notitle')}.png"
            new_path = out_dir / new_name
            Path(png).rename(new_path)
            generated.append(str(new_path))

    pdf.unlink(missing_ok=True)

    if grid:
        _apply_grid_overlay(generated)

    return {"preview_dir": str(out_dir), "files": generated, "output_path": str(out)}


def _extract_slide_titles(prs) -> dict[int, str]:
    """Extract sanitized slide titles from a Presentation object."""
    import re

    titles = {}
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title:
            title = slide.shapes.title.text.strip().replace("\n", " ")[:30]
        # Sanitise for filenames: drop reserved chars, collapse whitespace, trim.
        title = re.sub(r'[\\/:*?"<>|]', "", title)
        title = re.sub(r'\s+', "_", title).strip("_")
        titles[i] = title or "notitle"
    return titles


def _apply_grid_overlay(png_paths: list[str]) -> None:
    """Add coordinate grid overlay to PNG files."""
    from PIL import Image, ImageDraw, ImageFont

    color = (255, 0, 0, 128)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 24)
    except Exception:
        try:
            font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 24)
        except Exception:
            font = ImageFont.load_default()

    for png_path in png_paths:
        img = Image.open(png_path).convert("RGBA")
        w, h = img.size
        overlay = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        draw = ImageDraw.Draw(overlay)
        for pct in range(5, 100, 5):
            x, y = int(w * pct / 100), int(h * pct / 100)
            px_x, px_y = int(1920 * pct / 100), int(1080 * pct / 100)
            draw.line([(x, 0), (x, h)], fill=color, width=1)
            draw.line([(0, y), (w, y)], fill=color, width=1)
            if pct % 10 == 0:
                draw.text((x + 4, 4), f"{px_x}px ({pct}%)", fill=color, font=font)
                draw.text((4, y + 4), f"{px_y}px ({pct}%)", fill=color, font=font)
        Image.alpha_composite(img, overlay).convert("RGB").save(png_path)


def code_block(
    code: str,
    language: str = "python",
    theme: str = "dark",
    x: int = 0,
    y: int = 0,
    width: int = 800,
    height: int = 300,
    font_size: int = 12,
    show_label: bool = True,
) -> list[dict[str, Any]]:
    """Generate slide elements for a syntax-highlighted code block.

    Args:
        code: Source code text.
        language: Programming language for highlighting.
        theme: Color theme ("dark" or "light").
        x, y, width, height: Position and size in pixels.
        font_size: Code font size in pt.
        show_label: Show language label bar.

    Returns:
        List of element dicts for slide JSON.
    """
    from sdpm.builder.constants import CODE_COLORS
    from sdpm.utils.text import highlight_code

    colors = CODE_COLORS.get(theme, CODE_COLORS["dark"])
    bg = colors["background"]
    inverse_theme = "light" if theme == "dark" else "dark"
    inverse_bg = CODE_COLORS[inverse_theme]["background"]
    label_fg = "000000" if theme == "dark" else "FFFFFF"
    label_height = 22

    label_map = {"typescript": "TypeScript", "javascript": "JavaScript", "csharp": "C#", "cpp": "C++"}
    label_text = label_map.get(language, language.capitalize())

    elements: list[dict[str, Any]] = []
    if show_label:
        elements.append(
            {
                "type": "textbox",
                "x": x,
                "y": y,
                "width": width,
                "height": label_height,
                "fontSize": 8,
                "align": "left",
                "fill": inverse_bg,
                "text": f"{{{{#{label_fg}:{label_text}}}}}",
                "marginLeft": 50000,
                "marginTop": 0,
                "marginRight": 0,
                "marginBottom": 0,
                "autoWidth": True,
            }
        )
        code_y = y + label_height
        code_height = height - label_height
    else:
        code_y = y
        code_height = height

    spans = highlight_code(code, language, theme)
    elements.append(
        {
            "type": "textbox",
            "x": x,
            "y": code_y,
            "width": width,
            "height": code_height,
            "fontSize": font_size,
            "align": "left",
            "fill": bg,
            "text": spans,
        }
    )

    return elements
