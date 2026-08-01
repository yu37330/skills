# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Table element with CSS-style cascade styling."""
import sys

from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt


def _parse_background(value):
    """Parse background value. Returns (hex_color, alpha_pct) or (None, None).

    Supports:
      "#FF0000"         → ("#FF0000", None)
      "rgba(255,0,0,0.5)" → ("#FF0000", 50000)
      "none"            → ("none", None)
    """
    if not value or value == "none":
        return value, None
    if value.startswith("rgba("):
        inner = value[5:].rstrip(")")
        parts = [p.strip() for p in inner.split(",")]
        r, g, b = int(parts[0]), int(parts[1]), int(parts[2])
        a = float(parts[3])
        hex_color = f"#{r:02X}{g:02X}{b:02X}"
        alpha_pct = int(a * 100000)
        return hex_color, alpha_pct
    return value, None


def _resolve_cell_style(r, c, cell_val, has_headers, style, column_styles, cell_overrides):
    """Resolve final style for cell (r, c) via cascade.

    Cascade order: style(header/body/altRow/altCol) → columnStyles → cellOverrides → cell direct.
    """
    result = {}

    if style:
        if r == 0 and has_headers:
            result.update(style.get("header", {}))
        else:
            result.update(style.get("body", {}))
            data_ri = r - (1 if has_headers else 0)
            if data_ri % 2 == 1 and "altRow" in style:
                result.update(style["altRow"])
        if c % 2 == 1 and "altCol" in style:
            result.update(style["altCol"])

    if column_styles and str(c) in column_styles:
        result.update(column_styles[str(c)])

    if cell_overrides and f"{r},{c}" in cell_overrides:
        result.update(cell_overrides[f"{r},{c}"])

    # Cell direct props (highest priority)
    if isinstance(cell_val, dict):
        for k, v in cell_val.items():
            if k not in ("text", "gridSpan", "rowSpan", "merged",
                         "borders"):
                result[k] = v

    return result


class TableMixin:
    """Mixin providing table element methods."""

    def _add_table(self, slide, elem):
        from pptx.enum.text import PP_ALIGN
        from lxml import etree

        headers = elem.get("headers", [])
        rows = elem.get("rows", [])
        cols = len(headers) if headers else (len(rows[0]) if rows else 0)
        row_count = len(rows) + (1 if headers else 0)

        if cols == 0 or row_count == 0:
            print("Warning: table element without headers/rows skipped — element dropped from PPTX", file=sys.stderr)
            return

        x = self._px_to_emu(elem.get("x", 77))
        y = self._px_to_emu(elem.get("y", 270))
        width = self._px_to_emu(elem.get("width", 1766))
        height = self._px_to_emu(elem.get("height")) if elem.get("height") else Emu(row_count * 400000)

        tbl_shape = slide.shapes.add_table(row_count, cols, x, y, width, height)
        table = tbl_shape.table
        nsmap = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

        # Column widths
        col_widths = elem.get("colWidths")
        if col_widths:
            for i, w in enumerate(col_widths):
                if i < len(table.columns):
                    table.columns[i].width = self._px_to_emu(w)

        # Row heights
        row_heights = elem.get("rowHeights")
        if row_heights:
            for i, h in enumerate(row_heights):
                if i < len(table.rows):
                    table.rows[i].height = self._px_to_emu(h)

        # Clear table style — we handle all styling via cascade
        tbl_pr = table._tbl.find('a:tblPr', nsmap)
        if tbl_pr is not None:
            for attr in ['firstRow', 'lastRow', 'firstCol', 'lastCol', 'bandRow', 'bandCol']:
                if attr in tbl_pr.attrib:
                    del tbl_pr.attrib[attr]
            existing = tbl_pr.find('a:tableStyleId', nsmap)
            if existing is not None:
                tbl_pr.remove(existing)

        # Style inputs
        style = elem.get("style", {})
        column_styles = elem.get("columnStyles", {})
        cell_overrides = elem.get("cellOverrides", {})
        has_headers = bool(headers)
        global_border = style.get("border") if style else None

        # CSS convention: no border by default when style is specified
        clear_default_borders = bool(style)

        # Theme fallback when no style specified
        if not style:
            text_hex = self.theme_colors["text"].lstrip("#")
            bg_hex = self.theme_colors["background"].lstrip("#")
            br, bg, bb = int(bg_hex[:2], 16), int(bg_hex[2:4], 16), int(bg_hex[4:6], 16)
            shift = -12 if (0.299 * br + 0.587 * bg + 0.114 * bb) > 128 else 12
            alt = f"{max(0,min(255,br+shift)):02X}{max(0,min(255,bg+shift)):02X}{max(0,min(255,bb+shift)):02X}"
            style = {
                "header": {"background": f"#{text_hex}", "color": f"#{bg_hex}", "font-weight": "bold"},
                "body": {"background": f"#{bg_hex}", "color": f"#{text_hex}"},
                "altRow": {"background": f"#{alt}"},
            }

        # Merge cells first
        all_rows = [headers] + rows if headers else rows
        for ri, row_data in enumerate(all_rows):
            for ci, cell_val in enumerate(row_data):
                if not isinstance(cell_val, dict):
                    continue
                gs = cell_val.get("gridSpan", 1)
                rs = cell_val.get("rowSpan", 1)
                if gs > 1 or rs > 1:
                    try:
                        table.cell(ri, ci).merge(table.cell(ri + rs - 1, ci + gs - 1))
                    except Exception:
                        pass

        # Apply cells
        all_border_specs = {}
        for ri, row_data in enumerate(all_rows):
            for ci, cell_val in enumerate(row_data):
                # Skip merged-away cells
                if isinstance(cell_val, dict) and cell_val.get("merged"):
                    continue

                resolved = _resolve_cell_style(
                    ri, ci, cell_val, has_headers, style, column_styles, cell_overrides
                )
                text = cell_val if isinstance(cell_val, str) else (cell_val.get("text", "") if isinstance(cell_val, dict) else str(cell_val))
                cell = table.cell(ri, ci)

                # Ensure tcPr exists
                tc_pr = cell._tc.find(f'{{{nsmap["a"]}}}tcPr')
                if tc_pr is None:
                    tc_pr = etree.SubElement(cell._tc, f'{{{nsmap["a"]}}}tcPr')

                # Fill
                grad = resolved.get("gradient")
                if not grad and isinstance(cell_val, dict):
                    grad = cell_val.get("gradient")
                if grad:
                    from sdpm.builder.formatting import _build_grad_fill_element
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcPr.append(_build_grad_fill_element(grad))
                else:
                    bg_val = resolved.get("background")
                    hex_color, alpha_pct = _parse_background(bg_val)
                    if hex_color and hex_color != "none":
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(hex_color.lstrip('#'))
                        if alpha_pct is not None:
                            # Set fill transparency via XML alpha element
                            srgb_el = tc_pr.find(f'.//{{{nsmap["a"]}}}srgbClr')
                            if srgb_el is not None:
                                for old_alpha in srgb_el.findall(f'{{{nsmap["a"]}}}alpha'):
                                    srgb_el.remove(old_alpha)
                                alpha_el = etree.SubElement(srgb_el, f'{{{nsmap["a"]}}}alpha')
                                alpha_el.set('val', str(alpha_pct))
                    elif hex_color == "none":
                        cell.fill.background()

                # Vertical alignment
                _va_map = {"top": "t", "middle": "ctr", "bottom": "b"}
                va = resolved.get("vertical-align", "middle")
                tc_pr.set('anchor', _va_map.get(va, "ctr"))

                # Padding (cell inset)
                padding = resolved.get("padding")
                if not padding and isinstance(cell_val, dict):
                    padding = cell_val.get("padding")
                if padding:
                    for side, attr in [('left', 'marL'), ('right', 'marR'), ('top', 'marT'), ('bottom', 'marB')]:
                        if side in padding:
                            tc_pr.set(attr, str(self._px_to_emu(padding[side])))

                # Borders — global then cell-level override
                cell_borders = cell_val.get("borders") if isinstance(cell_val, dict) else None
                if not cell_borders and cell_overrides and f"{ri},{ci}" in cell_overrides:
                    cell_borders = cell_overrides[f"{ri},{ci}"].get("borders")
                border_spec = {}
                if clear_default_borders:
                    for side in ["left", "right", "top", "bottom"]:
                        border_spec[side] = {"fill": "none"}
                if global_border:
                    if any(k in global_border for k in ["left", "right", "top", "bottom"]):
                        for side, val in global_border.items():
                            border_spec[side] = val
                    else:
                        for side in ["left", "right", "top", "bottom"]:
                            border_spec[side] = dict(global_border)
                # Header-level border override
                if has_headers and ri == 0 and style.get("header", {}).get("border"):
                    hdr_border = style["header"]["border"]
                    if any(k in hdr_border for k in ["left", "right", "top", "bottom"]):
                        for side, val in hdr_border.items():
                            border_spec[side] = val
                    else:
                        for side in ["left", "right", "top", "bottom"]:
                            border_spec[side] = dict(hdr_border)
                if cell_borders:
                    for side, val in cell_borders.items():
                        border_spec[side] = val

                all_border_specs[(ri, ci)] = border_spec

                # Text
                para = cell.text_frame.paragraphs[0]
                text_color = None
                color_val = resolved.get("color")
                if color_val:
                    text_color = RGBColor.from_string(color_val.lstrip('#'))

                align = resolved.get("text-align")
                align_val = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align) if align else None
                if align_val is not None:
                    para.alignment = align_val

                cell_paras = cell_val.get("paragraphs") if isinstance(cell_val, dict) else None
                if cell_paras:
                    # Paragraph list with optional bullets: one a:p per item.
                    tf = cell.text_frame
                    a_ns = nsmap["a"]
                    for pi, pd in enumerate(cell_paras):
                        if isinstance(pd, str):
                            pd = {"text": pd}
                        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                        if align_val is not None:
                            p.alignment = align_val
                        self._apply_styled_text(p, str(pd.get("text", "")), default_color=text_color, no_default_color=text_color is None)
                        pPr = p._element.get_or_add_pPr()
                        bu = pd.get("bullet")
                        if bu:
                            char = bu if isinstance(bu, str) else "•"
                            bu_font = etree.SubElement(pPr, f'{{{a_ns}}}buFont')
                            bu_font.set('typeface', 'Arial')
                            bu_char = etree.SubElement(pPr, f'{{{a_ns}}}buChar')
                            bu_char.set('char', char)
                            # Hanging indent so wrapped lines align after the
                            # bullet — explicit marL/indent (EMU) win.
                            pPr.set('marL', str(pd.get('marL', 171450)))
                            pPr.set('indent', str(pd.get('indent', -171450)))
                        else:
                            etree.SubElement(pPr, f'{{{a_ns}}}buNone')
                            if pd.get('marL') is not None:
                                pPr.set('marL', str(pd['marL']))
                            if pd.get('indent') is not None:
                                pPr.set('indent', str(pd['indent']))
                    cell_runs = [r for p in cell.text_frame.paragraphs for r in p.runs]
                else:
                    self._apply_styled_text(para, str(text), default_color=text_color, no_default_color=text_color is None)
                    cell_runs = para.runs

                # Font properties on all runs
                font_weight = resolved.get("font-weight")
                font_style = resolved.get("font-style")
                text_decoration = resolved.get("text-decoration")
                font_size = resolved.get("font-size")
                for run in cell_runs:
                    if font_weight == "bold":
                        run.font.bold = True
                    if font_style == "italic":
                        run.font.italic = True
                    if text_decoration == "underline":
                        run.font.underline = True
                    if font_size and run.font.size is None:
                        run.font.size = Pt(font_size)

        # Propagate borders to adjacent cells (CSS convention: one-side spec is enough)
        opposite = {"top": "bottom", "bottom": "top", "left": "right", "right": "left"}
        row_delta = {"top": -1, "bottom": 1, "left": 0, "right": 0}
        col_delta = {"top": 0, "bottom": 0, "left": -1, "right": 1}
        for (ri, ci), bspec in list(all_border_specs.items()):
            for side, bdr in bspec.items():
                if bdr.get("fill") == "none":
                    continue
                nr, nc = ri + row_delta[side], ci + col_delta[side]
                if (nr, nc) not in all_border_specs:
                    continue
                opp = opposite[side]
                neighbor = all_border_specs[(nr, nc)]
                if neighbor.get(opp, {}).get("fill") == "none":
                    neighbor[opp] = dict(bdr)

        # Write borders to XML
        num_rows = len(all_rows)
        num_cols = len(all_rows[0]) if all_rows else 0
        for ri in range(num_rows):
            for ci in range(num_cols):
                bspec = all_border_specs.get((ri, ci))
                if not bspec:
                    continue
                cell = table.cell(ri, ci)
                tc_pr = cell._tc.find(f'{{{nsmap["a"]}}}tcPr')
                if tc_pr is None:
                    tc_pr = etree.SubElement(cell._tc, f'{{{nsmap["a"]}}}tcPr')
                tag_map = {"left": "lnL", "right": "lnR", "top": "lnT", "bottom": "lnB"}
                ref_el = tc_pr.find(f'{{{nsmap["a"]}}}solidFill')
                if ref_el is None:
                    ref_el = tc_pr.find(f'{{{nsmap["a"]}}}noFill')
                if ref_el is None:
                    ref_el = tc_pr.find(f'{{{nsmap["a"]}}}gradFill')
                for side, bdr in bspec.items():
                    tag = tag_map.get(side)
                    if not tag:
                        continue
                    existing = tc_pr.find(f'{{{nsmap["a"]}}}{tag}')
                    if existing is not None:
                        tc_pr.remove(existing)
                    ln = etree.Element(f'{{{nsmap["a"]}}}{tag}')
                    if bdr.get("width"):
                        ln.set('w', str(int(bdr["width"] * 12700)))
                    if bdr.get("fill") == "none":
                        etree.SubElement(ln, f'{{{nsmap["a"]}}}noFill')
                    elif bdr.get("color"):
                        sf = etree.SubElement(ln, f'{{{nsmap["a"]}}}solidFill')
                        srgb = etree.SubElement(sf, f'{{{nsmap["a"]}}}srgbClr')
                        srgb.set('val', bdr["color"].lstrip('#'))
                    if ref_el is not None:
                        tc_pr.insert(list(tc_pr).index(ref_el), ln)
                    else:
                        tc_pr.append(ln)

        # Apply effects (shadow, glow, etc.) to the table shape
        from sdpm.utils.effects import apply_effects
        apply_effects(tbl_shape._element, elem, self.EMU_PER_PX)
