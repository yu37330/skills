# Copyright Amazon.com, Inc. or its affiliates. All Rights Reserved.
# SPDX-License-Identifier: MIT-0
"""Chart extraction."""
from pptx.enum.chart import XL_CHART_TYPE

from .constants import _base_element, _hex
from .color import _resolve_scheme_color

def extract_chart_element(shape, theme_colors=None, color_mapping=None):
    """Extract chart element from a GraphicFrame containing a chart."""
    chart = shape.chart
    elem = _base_element(shape, "chart")

    # Map chart type
    ct = chart.chart_type
    type_map = {
        XL_CHART_TYPE.COLUMN_CLUSTERED: ("bar", False, False),
        XL_CHART_TYPE.COLUMN_STACKED: ("bar", True, False),
        XL_CHART_TYPE.BAR_CLUSTERED: ("bar", False, True),
        XL_CHART_TYPE.BAR_STACKED: ("bar", True, True),
        XL_CHART_TYPE.LINE: ("line", False, False),
        XL_CHART_TYPE.LINE_MARKERS: ("line", False, False),
        XL_CHART_TYPE.PIE: ("pie", False, False),
        XL_CHART_TYPE.DOUGHNUT: ("donut", False, False),
    }
    chart_type, stacked, horizontal = type_map.get(ct, ("bar", False, False))
    elem["chartType"] = chart_type
    if stacked:
        elem["stacked"] = True
    if horizontal:
        elem["horizontal"] = True
    if ct == XL_CHART_TYPE.LINE:
        elem["markers"] = False

    # Title
    if chart.has_title:
        try:
            elem["title"] = chart.chart_title.text_frame.text
            # Extract title font color if explicit
            runs = chart.chart_title.text_frame.paragraphs[0].runs
            if runs:
                from pptx.oxml.ns import qn as _qn
                rPr = runs[0]._r.find(_qn('a:rPr'))
                if rPr is not None:
                    sf = rPr.find(_qn('a:solidFill'))
                    if sf is not None:
                        srgb = sf.find(_qn('a:srgbClr'))
                        if srgb is not None:
                            elem["titleFontColor"] = _hex(srgb)
                    sz = rPr.get('sz')
                    if sz:
                        elem["titleFontSize"] = int(sz) // 100
        except Exception:
            pass

    # Categories and series
    plot = chart.plots[0]
    try:
        cats = [str(c) for c in plot.categories]
        elem["categories"] = cats
    except Exception:
        elem["categories"] = []

    series_list = []
    for series in plot.series:
        s = {"name": ""}
        try:
            if hasattr(series, 'tx') and hasattr(series.tx, 'text'):
                s["name"] = str(series.tx.text)
            elif hasattr(series, '_element'):
                ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
                tx = series._element.find('.//c:tx//c:v', ns)
                if tx is not None and tx.text:
                    s["name"] = tx.text
        except Exception:
            pass
        try:
            s["values"] = list(series.values)
        except Exception:
            s["values"] = []
        # Extract color (solid or gradient fallback)
        try:
            if chart_type in ("bar",):
                rgb = series.format.fill.fore_color.rgb
                s["color"] = f"#{rgb}"
            elif chart_type == "line":
                rgb = series.format.line.color.rgb
                s["color"] = f"#{rgb}"
        except Exception:
            # Try gradient: use last stop color as representative
            try:
                from pptx.oxml.ns import qn as _qn
                ser_el = series._element
                gradFill = ser_el.find(f'.//{_qn("a:gradFill")}')
                if gradFill is not None:
                    stops = gradFill.findall(f'.//{_qn("a:gs")}')
                    if stops:
                        last = stops[-1]
                        srgb = last.find(f'{_qn("a:srgbClr")}')
                        scheme = last.find(f'{_qn("a:schemeClr")}')
                        if srgb is not None:
                            s["color"] = _hex(srgb)
                        elif scheme is not None:
                            resolved = _resolve_scheme_color(scheme.get('val'), theme_colors, color_mapping)
                            if resolved:
                                s["color"] = resolved
            except Exception:
                pass
        series_list.append(s)
    elem["series"] = series_list
    if not series_list:
        # Empty chart (decorative) — still preserve XML for roundtrip
        from lxml import etree as _et
        elem["_chartXml"] = _et.tostring(chart._chartSpace, encoding='unicode')
        for rel in chart.part.rels.values():
            if rel.reltype.endswith('/chartUserShapes'):
                elem["_chartUserShapes"] = rel.target_part.blob.decode('utf-8')
                break
        elem["series"] = [{"name": "", "values": [1]}]  # dummy for builder
        return elem

    # Per-point colors (dPt)
    try:
        ns_c = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
                'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for si, ser_el in enumerate(chart._chartSpace.findall('.//c:ser', ns_c)):
            point_colors = {}
            point_xml = {}
            for dPt in ser_el.findall('c:dPt', ns_c):
                pt_idx = int(dPt.find('c:idx', ns_c).get('val'))
                spPr = dPt.find('c:spPr', ns_c)
                if spPr is not None:
                    solid = spPr.find('a:solidFill', ns_c)
                    grad = spPr.find('a:gradFill', ns_c)
                    if solid is not None:
                        srgb = solid.find('a:srgbClr', ns_c)
                        scheme = solid.find('a:schemeClr', ns_c)
                        if srgb is not None:
                            point_colors[str(pt_idx)] = _hex(srgb)
                        elif scheme is not None:
                            resolved = _resolve_scheme_color(scheme.get('val'), theme_colors, color_mapping)
                            if resolved:
                                point_colors[str(pt_idx)] = resolved
                    elif grad is not None:
                        # Preserve gradient as XML
                        from lxml import etree as _et
                        point_xml[str(pt_idx)] = _et.tostring(spPr, encoding='unicode')
            if point_colors and si < len(series_list):
                series_list[si]["pointColors"] = point_colors
            if point_xml and si < len(series_list):
                series_list[si]["_pointXml"] = point_xml
    except Exception:
        pass
    try:
        val_axis = chart.value_axis
        axis_info = {}
        if val_axis.maximum_scale is not None:
            axis_info["max"] = val_axis.maximum_scale
        if val_axis.minimum_scale is not None:
            axis_info["min"] = val_axis.minimum_scale
        if val_axis.major_unit is not None:
            axis_info["majorUnit"] = val_axis.major_unit
        if not val_axis.has_major_gridlines:
            axis_info["gridlines"] = False
        if str(val_axis.major_tick_mark) == 'NONE (-4142)':
            axis_info["tickMark"] = "none"
        # Axis line
        ns_c = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
                'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        valAx_el = chart._chartSpace.find('.//c:valAx', ns_c)
        if valAx_el is not None:
            val_spPr = valAx_el.find('c:spPr', ns_c)
            if val_spPr is not None:
                ln = val_spPr.find('a:ln', ns_c)
                if ln is not None and ln.find('a:noFill', ns_c) is not None:
                    axis_info["line"] = "none"
        if axis_info:
            elem["valueAxis"] = axis_info
    except Exception:
        pass

    # Category axis
    try:
        cat_axis = chart.category_axis
        cat_info = {}
        if not cat_axis.has_major_gridlines:
            cat_info["gridlines"] = False
        if str(cat_axis.major_tick_mark) == 'NONE (-4142)':
            cat_info["tickMark"] = "none"
        if cat_info:
            elem["categoryAxis"] = cat_info
    except Exception:
        pass

    # Legend
    elem["legend"] = chart.has_legend

    # Bar gap/overlap
    if chart_type == "bar":
        try:
            ns_c3 = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            barChart = chart._chartSpace.find('.//c:barChart', ns_c3)
            if barChart is not None:
                gw = barChart.find('c:gapWidth', ns_c3)
                if gw is not None:
                    elem["gapWidth"] = int(gw.get('val'))
                ov = barChart.find('c:overlap', ns_c3)
                if ov is not None:
                    elem["overlap"] = int(ov.get('val'))
        except Exception:
            pass

    # Chart font size (from axis/data label defRPr)
    try:
        ns_c2 = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
                 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        for tag in ['c:catAx', 'c:valAx']:
            ax = chart._chartSpace.find(f'.//{tag}', ns_c2)
            if ax is not None:
                defRPr = ax.find('.//c:txPr//a:defRPr', ns_c2)
                if defRPr is not None and defRPr.get('sz'):
                    elem["_chartFontSize"] = int(defRPr.get('sz'))
                    break
    except Exception:
        pass

    # Data labels
    if plot.has_data_labels:
        # Check if any label type is actually shown
        ns_c = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        dLbls = chart._chartSpace.find('.//c:dLbls', ns_c)
        show_any = False
        if dLbls is not None:
            for tag in ('showVal', 'showCatName', 'showSerName', 'showPercent'):
                el = dLbls.find(f'c:{tag}', ns_c)
                if el is not None and el.get('val') == '1':
                    show_any = True
        if show_any:
            elem["dataLabels"] = True
        try:
            nf = plot.data_labels.number_format
            if nf:
                elem["numberFormat"] = nf
        except Exception:
            pass

    # Donut hole size
    if chart_type == "donut":
        try:
            ns_c = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            hs = chart._chartSpace.find('.//c:holeSize', ns_c)
            if hs is not None:
                elem["holeSize"] = int(hs.get('val'))
        except Exception:
            pass

    # Preserve full chart XML for perfect roundtrip
    try:
        from lxml import etree as _et
        elem["_chartXml"] = _et.tostring(chart._chartSpace, encoding='unicode')
        # Save chartUserShapes (text overlays on chart)
        for rel in chart.part.rels.values():
            if rel.reltype.endswith('/chartUserShapes'):
                elem["_chartUserShapes"] = rel.target_part.blob.decode('utf-8')
                break
    except Exception:
        pass

    return elem
